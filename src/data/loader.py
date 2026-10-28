"""Safe reading of uploaded CSV and Excel files.

The file is only ever parsed as data: nothing in it is executed. Excel files
are read with openpyxl, which ignores macros and uses stored cell values
rather than recalculating formulas. Only .csv and .xlsx are accepted.
"""

import csv
import io
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ALLOWED_EXTENSIONS = (".csv", ".xlsx")
MAX_FILE_MB = 50
MAX_UNCOMPRESSED_MB = 500   # guards against "zip bombs" disguised as .xlsx
MAX_ROWS = 2_000_000
XLSX_SIGNATURE = b"PK\x03\x04"  # every .xlsx is a zip archive
CSV_ENCODINGS = ("utf-8-sig", "cp1252", "latin-1")
CSV_DELIMITERS = ",;\t|"
SHEET_COLUMN = "SourceSheet"


class UploadError(Exception):
    """The file cannot be accepted; the message is shown to the user."""


@dataclass
class LoadedFile:
    name: str
    data: pd.DataFrame
    notes: list[str] = field(default_factory=list)


def validate_upload(name: str, content: bytes, max_mb: float = MAX_FILE_MB) -> str:
    """Check name, size and file signature. Returns the lower-case extension."""
    extension = Path(name).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise UploadError(
            f"'{extension or 'no extension'}' files are not supported. Please upload a .csv or .xlsx file."
        )
    if not content:
        raise UploadError("The file is empty.")
    size_mb = len(content) / 1_048_576
    if size_mb > max_mb:
        raise UploadError(f"The file is {size_mb:.1f} MB; the limit is {max_mb} MB.")

    if extension == ".xlsx":
        if not content.startswith(XLSX_SIGNATURE):
            raise UploadError("This file is named .xlsx but is not a valid Excel workbook.")
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                unpacked = sum(info.file_size for info in archive.infolist())
        except zipfile.BadZipFile as exc:
            raise UploadError("The Excel file is damaged and cannot be opened.") from exc
        if unpacked / 1_048_576 > MAX_UNCOMPRESSED_MB:
            raise UploadError("The Excel file expands to more than "
                              f"{MAX_UNCOMPRESSED_MB} MB and was rejected for safety.")
    elif b"\x00" in content[:65_536]:
        raise UploadError("This does not look like a text CSV file (it contains binary data).")
    return extension


def list_sheets(content: bytes) -> list[str]:
    workbook = load_workbook(io.BytesIO(content), read_only=True)
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


def _detect_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=CSV_DELIMITERS).delimiter
    except csv.Error:
        return ","


def _read_csv(content: bytes) -> tuple[pd.DataFrame, list[str]]:
    for encoding in CSV_ENCODINGS:
        try:
            text_sample = content[:65_536].decode(encoding)
        except UnicodeDecodeError:
            continue
        delimiter = _detect_delimiter(text_sample)
        try:
            df = pd.read_csv(io.BytesIO(content), encoding=encoding, sep=delimiter, low_memory=False)
        except UnicodeDecodeError:
            continue
        except (pd.errors.ParserError, pd.errors.EmptyDataError) as exc:
            raise UploadError(f"The CSV file could not be read: {exc}") from exc
        notes = [f"Read as CSV (encoding {encoding}, delimiter {delimiter!r})."]
        return df, notes
    raise UploadError("The CSV file's text encoding could not be recognised.")


def _read_excel(content: bytes, sheets: list[str] | None) -> tuple[pd.DataFrame, list[str]]:
    available = list_sheets(content)
    chosen = sheets or available[:1]
    unknown = [s for s in chosen if s not in available]
    if unknown:
        raise UploadError(f"Sheet(s) not found: {unknown}")
    frames = pd.read_excel(io.BytesIO(content), sheet_name=chosen, engine="openpyxl")
    if len(chosen) == 1:
        return frames[chosen[0]], [f"Read sheet '{chosen[0]}'."]

    first_columns = list(frames[chosen[0]].columns)
    for name in chosen[1:]:
        if list(frames[name].columns) != first_columns:
            raise UploadError(f"Sheet '{name}' has different columns from '{chosen[0]}', "
                              "so the sheets cannot be combined.")
    combined = pd.concat([df.assign(**{SHEET_COLUMN: name}) for name, df in frames.items()],
                         ignore_index=True)
    return combined, [f"Combined sheets {chosen} (column '{SHEET_COLUMN}' records the source sheet)."]


def read_upload(name: str, content: bytes, sheets: list[str] | None = None,
                max_rows: int = MAX_ROWS) -> LoadedFile:
    """Validate and parse an uploaded file into a DataFrame."""
    extension = validate_upload(name, content)
    try:
        if extension == ".csv":
            df, notes = _read_csv(content)
        else:
            df, notes = _read_excel(content, sheets)
    except UploadError:
        raise
    except Exception as exc:  # any other parsing failure becomes a friendly message
        raise UploadError(f"The file could not be read ({type(exc).__name__}).") from exc

    if df.empty or df.shape[1] < 2:
        raise UploadError("The file has no data rows or fewer than two columns.")
    if len(df) > max_rows:
        raise UploadError(f"The file has {len(df):,} rows; the limit is {max_rows:,}.")
    unnamed = [c for c in df.columns if str(c).startswith("Unnamed:")]
    if unnamed:
        notes.append(f"{len(unnamed)} column(s) have no header name: {unnamed}.")
    return LoadedFile(name=name, data=df, notes=notes)
